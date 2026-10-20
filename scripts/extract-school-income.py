"""Offline, reviewed contract-income layout only. No network or formula execution.

Uses the existing local openpyxl runtime. Originals and output MUST stay outside
the public repository. Review original contents before using the generated audit.
"""
import argparse
import hashlib
import json
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile
import re

import openpyxl
from openpyxl.utils import get_column_letter as col


def extract(path, library_id, library_version):
    raw = Path(path).read_bytes()
    if not 0 < len(raw) <= 5 * 1024 * 1024:
        raise ValueError("Unsupported source size")
    with ZipFile(path) as archive:
        parts = archive.infolist()
        if len(parts) > 300 or sum(p.file_size for p in parts) > 20 * 1024 * 1024:
            raise ValueError("Workbook archive exceeds limits")
        if len({p.filename for p in parts}) != len(parts):
            raise ValueError("Duplicate workbook parts")
        if any(any(marker in p.filename.lower() for marker in ["externallinks", "vbaproject", "embeddings"]) for p in parts):
            raise ValueError("External or embedded workbook content requires review")
        for p in parts:
            if p.filename.endswith((".xml", ".rels")):
                data = archive.read(p)
                if b"<!DOCTYPE" in data or b"<!ENTITY" in data or b'TargetMode="External"' in data:
                    raise ValueError("Unsupported external XML content")
    formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    if not 1 <= len(formulas.worksheets) <= 20:
        raise ValueError("Unsupported sheet count")
    result = {
        "version": "school-income-audit-v1",
        "source": {"sha256": hashlib.sha256(raw).hexdigest(), "filename": Path(path).name,
                   "byteSize": len(raw), "libraryId": library_id, "libraryVersion": library_version},
        "currency": None, "accountingBasis": None, "hoursBasis": None, "mappings": None,
        "sheets": [],
    }
    for sheet in formulas:
        if not re.fullmatch(r"20\d{2}", sheet.title) or sheet.sheet_state != "visible" or sheet.max_row > 1100 or sheet.max_column > 40:
            raise ValueError("Unsupported sheet layout")
        values = cached[sheet.title]
        for row in sheet:
            for c in row:
                if c.comment or c.hyperlink:
                    raise ValueError("Comments or hyperlinks require separate review")
                if isinstance(c.value, str) and re.search(r"\b\d{3}-\d{2}-\d{4}\b|bank\s*account|routing\s*number|student\s*(name|id)|patient\s*(name|id)", c.value, re.I):
                    raise ValueError("Sensitive-field marker requires review")
        if any(d.hidden for d in sheet.row_dimensions.values()) or any(d.hidden for d in sheet.column_dimensions.values()):
            raise ValueError("Hidden rows or columns require review")

        def cell(row, column, precision=2):
            c, v = sheet.cell(row, column), values.cell(row, column)
            formula = c.value if c.data_type == "f" else None
            value = v.value
            if v.data_type == "e":
                return dict(ref=c.coordinate, state="error", value=str(value), formula=formula)
            if value is None:
                if formula:
                    raise ValueError(f"Formula has no cached result: {sheet.title}!{c.coordinate}")
                return dict(ref=c.coordinate, state="blank", value=None, formula=None)
            if isinstance(value, str) and value.strip() == "-" and formula is None:
                return dict(ref=c.coordinate, state="dash", value="-", formula=None)
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"Unsupported source value: {sheet.title}!{c.coordinate}")
            number = Decimal(str(value))
            if not number.is_finite() or number != number.quantize(Decimal(10) ** -precision):
                raise ValueError(f"Unsupported source precision: {sheet.title}!{c.coordinate}")
            return dict(ref=c.coordinate, state="number", value=format(number, "f"), formula=formula)

        leading = sheet["B3"].value == "Dec hours" and sheet["C3"].value == "Dec"
        income_columns = [3, 5, 7, 9, 11, 13, 15, 18, 20, 22, 24, 26] if leading else list(range(3, 24, 2))
        months = ([None] if leading else []) + [1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12]
        standard_headers = ["Jan", "Feb", "March", "April", "May", "June", "July/Aug", "Sep", "Oct", "Nov", "Dec"]
        if leading:
            expected = {"B3": "Dec hours", "C3": "Dec", "E3": "Jan", "G3": "Feb", "I3": "March",
                        "K3": "April", "M3": "May", "O3": "June", "P3": "July", "R3": "Aug",
                        "S3": "Sep", "T3": "Sep", "U3": "Oct", "W3": "Nov", "Y3": "Dec"}
        else:
            expected = {f"{col(c)}3": label for c, label in zip(income_columns, standard_headers)}
        actual = {c.coordinate: c.value for c in sheet[3] if c.value is not None}
        if actual != expected:
            raise ValueError("Unrecognized period headers")
        periods = []
        for c, month in zip(income_columns, months):
            periods.append(dict(id="unassigned-december" if month is None else "07-08" if month == 7 else f"{month:02}",
                                label="December · year unconfirmed" if month is None else "July / August combined" if month == 7 else standard_headers[[1,2,3,4,5,6,7,9,10,11,12].index(month)],
                                start=month, end=8 if month == 7 else month, hoursColumn=col(c-1), incomeColumn=col(c)))
        detail_rows = [r for r in range(4, sheet.max_row + 1)
                       if isinstance(sheet.cell(r, 1).value, str)
                       and sheet.cell(r, 1).value.strip().lower() not in ("hours", "income", "")]
        if not detail_rows:
            raise ValueError("No school rows")
        for r in range(4, max(detail_rows) + 1):
            if r not in detail_rows and any(sheet.cell(r, c).value is not None for c in range(1, sheet.max_column + 1)):
                raise ValueError("Unlabeled or unsupported row inside school detail")
        rows = []
        for r in range(4, sheet.max_row + 1):
            label = sheet.cell(r, 1).value
            if isinstance(label, str) and label.strip().lower() not in ("hours", "income", ""):
                if sheet.cell(r, 1).data_type == "f":
                    raise ValueError("Formula school label")
                # An unexpected populated column must never disappear from extraction.
                allowed = {1} | {c for income in income_columns for c in (income-1, income)}
                if any(sheet.cell(r, c).value is not None for c in range(1, sheet.max_column + 1) if c not in allowed):
                    raise ValueError("Unrecognized school detail outside the reviewed column layout")
                rows.append(dict(row=r, label=label, values=[dict(hours=cell(r,c-1,4), income=cell(r,c)) for c in income_columns]))
        if not rows:
            raise ValueError("No school rows")
        controls = []
        for row in sheet.iter_rows(min_row=max(r["row"] for r in rows)+1):
            for c in row:
                if c.data_type == "f":
                    controls.append(cell(c.row, c.column, 8))
        notes = ["Source currency, accounting basis, hours meaning and school/location/provider mappings are unconfirmed.",
                 "Combined July/August is one interval. Blank, dash and error cells are not zero. Numeric subtotals do not establish complete period coverage."]
        if leading:
            notes.append("Leading December has an unconfirmed year and is excluded from selected calendar periods.")
        by_label = {}
        for row in rows:
            by_label.setdefault(row["label"].strip().casefold(), []).append(row["row"])
        for duplicate in by_label.values():
            if len(duplicate) > 1:
                notes.append(f"Repeated label on source rows {', '.join(map(str,duplicate))}; retained separately pending mapping review.")
        # Compare source SUM controls to numeric leaf cells without evaluating Excel.
        for control in controls:
            match = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):\1(\d+)\)", control["formula"] or "", re.I)
            if not match or control["state"] != "number":
                continue
            column = match[1]
            if column not in [p["incomeColumn"] for p in periods]:
                continue
            i = [p["incomeColumn"] for p in periods].index(column)
            observed = sum((Decimal(r["values"][i]["income"]["value"]) for r in rows if r["values"][i]["income"]["state"] == "number"), Decimal(0))
            delta = observed - Decimal(control["value"])
            if delta:
                notes.append(f"{control['ref']} cached source control differs from numeric detail by {delta:f}; formula {control['formula']}. Detail subtotal is retained separately.")
        if leading:
            notes.append("Annual source control excludes later period columns; do not use it as a complete annual total.")
        numeric = [v["income"] for r in rows for v in r["values"] if v["income"]["state"] == "number"]
        pairs = sum(v["hours"]["state"] != "blank" or v["income"]["state"] != "blank" for r in rows for v in r["values"])
        total = sum((Decimal(c["value"]) for c in numeric), Decimal(0))
        result["sheets"].append(dict(sheet=sheet.title, year=int(sheet.title), periods=periods, rows=rows, controls=controls, notes=notes,
                                     numericCells=len(numeric), populatedPairs=pairs, rawHundredths=str(int(total*100))))
    # Flag unchanged October activity in adjacent exports for review, without dropping it.
    ordered = sorted(result["sheets"], key=lambda s: s["year"])
    for previous, current in zip(ordered, ordered[1:]):
        if current["year"] != previous["year"] + 1:
            continue
        p_index = next(i for i,p in enumerate(previous["periods"]) if p["id"] == "10")
        c_index = next(i for i,p in enumerate(current["periods"]) if p["id"] == "10")
        for row in current["rows"]:
            pair = row["values"][c_index]
            if pair["income"]["state"] != "number" or Decimal(pair["income"]["value"]) == 0:
                continue
            matches = [r for r in previous["rows"] if r["label"].strip().casefold() == row["label"].strip().casefold()
                       and all(r["values"][p_index][k]["value"] == pair[k]["value"] and r["values"][p_index][k]["state"] == pair[k]["state"] for k in ("hours", "income"))]
            if matches:
                current["notes"].append(f"October source row {row['row']} matches hours/income in {previous['sheet']} row(s) {', '.join(str(r['row']) for r in matches)}; confirm activity rather than assuming a carry-forward error.")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--library-id", required=True)
    parser.add_argument("--library-version", type=int, required=True)
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    for path in [args.input, args.output]:
        if Path(path).resolve().is_relative_to(repo):
            raise ValueError("Private files must remain outside the public repository")
    audit = extract(args.input, args.library_id, args.library_version)
    with open(args.output, "x") as handle:
        import os
        os.chmod(args.output, 0o600)
        json.dump(audit, handle, separators=(",", ":"))
    print(json.dumps({"sheets": len(audit["sheets"]), "schoolRows": sum(len(s["rows"]) for s in audit["sheets"]),
                      "numericCells": sum(s["numericCells"] for s in audit["sheets"]), "populatedPairs": sum(s["populatedPairs"] for s in audit["sheets"])}))
